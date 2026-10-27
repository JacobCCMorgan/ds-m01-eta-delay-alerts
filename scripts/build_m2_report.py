"""Author the four-page M2 report from completed execution receipts.

Run only after the document-operation marker, then render and visually review
every page with the document skill. This script does not perform data builds.
"""
from pathlib import Path
import argparse
import json
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASENAME = 'DS-M01_M2_Results_20261027_v0_2'
ROLE_NAMES = {'train': '訓練', 'calibration': '校準', 'test': '測試', 'excluded': '排除'}
COHORT_NAMES = {
    'seen_vehicle_and_route': '已見車輛與路線',
    'unseen_vehicle_only': '未見車輛',
    'unseen_route_only': '未見路線',
    'unseen_vehicle_and_route': '未見車輛及路線',
}


def load_receipts(root):
    reports = root / 'reports'
    names = {
        'summary': 'm2_summary.json', 'quality': 'quality_report.json',
        'readiness': 'm2_readiness.json', 'rebuild': 'rebuild_report.json',
        'http': 'last_http-check.json', 'runtime': 'runtime_profile.json',
        'compatibility': 'pipeline_m1_compatibility.json', 'postgres': 'postgres_check.json',
    }
    data = {key: json.loads((reports / filename).read_text()) for key, filename in names.items()}
    tests = (reports / 'unittest_results.txt').read_text()
    count = re.search(r'Ran\s+(\d+)\s+tests?\s+in', tests)
    if not count or not re.search(r'^OK\s*$', tests, re.MULTILINE):
        raise ValueError('A complete passing unittest receipt is required')
    data['test_count'] = int(count.group(1))
    required = {
        'warehouse quality': data['quality'].get('passed') is True,
        'M2 technical readiness': data['readiness'].get('technical_checks_passed') is True,
        'fresh rebuild': data['rebuild'].get('status') == 'passed',
        'HTTP': data['http'].get('status') == 'passed',
        'M1 pipeline compatibility': data['compatibility'].get('status') == 'passed',
        'M1 business-row continuity': data['readiness'].get('m1_continuity', {}).get('status') == 'passed',
    }
    failures = [name for name, passed in required.items() if not passed]
    if failures:
        raise ValueError('Cannot author success report before passing: ' + ', '.join(failures))
    if data['summary']['trip_count'] != data['quality']['counts']['eta_trip']:
        raise ValueError('Summary and warehouse trip counts disagree')
    if data['summary']['snapshot_count'] != data['quality']['counts']['eta_snapshot']:
        raise ValueError('Summary and warehouse snapshot counts disagree')
    return data


def draw_chart(summary, path):
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10,
                         'axes.spines.top': False, 'axes.spines.right': False})
    figure, axes = plt.subplots(1, 2, figsize=(10.5, 3.4), layout='constrained')
    panels = [
        (['Eligible', 'Not mature', 'Cancelled', 'Already arrived'],
         [summary['label_eligibility'].get(k, 0) for k in
          ['eligible', 'not_mature', 'cancelled', 'already_arrived']], 'Target availability'),
        (['Valid', 'Stale', 'No position'],
         [summary['position_quality'].get(k, 0) for k in ['valid', 'low_confidence', 'no_position']],
         'Position quality'),
    ]
    for axis, (labels, values, title) in zip(axes, panels):
        axis.barh(labels, values, color=['#286B86', '#C58D36', '#82919A', '#AA6254'][:len(values)])
        axis.invert_yaxis()
        axis.set_title(title, loc='left', fontsize=12, weight='bold')
        axis.set_xlabel('Snapshots')
        maximum = max(values) or 1
        axis.set_xlim(0, maximum * 1.27)
        for index, value in enumerate(values):
            axis.text(value + maximum * .025, index, f'{value:,}', va='center', fontsize=10)
    figure.savefig(path, dpi=190)
    plt.close(figure)


class Report:
    def __init__(self):
        self.doc = Document()
        # Remove the template's inherited title rule; headings use whitespace.
        for border in list(self.doc.styles.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
        self.text = []
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.62)
        section.left_margin = section.right_margin = Inches(.7)
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
            style = self.doc.styles[name]
            style.font.name = 'Noto Sans CJK TC'
            style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK TC')
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(10.3)
            style.paragraph_format.line_spacing = 1.1
            style.paragraph_format.space_after = Pt(6)
        for name, size in [('Title', 22), ('Heading 1', 17), ('Heading 2', 12), ('Caption', 8.5)]:
            self.doc.styles[name].font.size = Pt(size)
        self.doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
        self.doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
        self.doc.styles['Heading 2'].font.bold = True
        self.doc.styles['Caption'].font.color.rgb = RGBColor.from_string('52616D')
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('DS-M01 M2  |  ').font.size = Pt(8)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)

    def p(self, text, style='Normal'):
        self.doc.add_paragraph(text, style)
        self.text.append(text)

    def page(self, title):
        self.doc.add_page_break()
        self.p(title, 'Heading 1')

    def table(self, headers, rows, widths, numeric=()):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for column, value in enumerate(headers):
            table.rows[0].cells[column].text = str(value)
        for values in rows:
            cells = table.add_row().cells
            for column, value in enumerate(values):
                cells[column].text = str(value)
        for index, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if index == 0:
                props.append(OxmlElement('w:tblHeader'))
            for column, cell in enumerate(row.cells):
                cell.width = Inches(widths[column])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell_props = cell._tc.get_or_add_tcPr()
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '224C65' if index == 0 else ('F1F5F8' if index % 2 else 'FFFFFF'))
                cell_props.append(shading)
                borders = OxmlElement('w:tcBorders')
                for edge in ('top', 'left', 'bottom', 'right'):
                    border = OxmlElement('w:' + edge)
                    for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                        border.set(qn('w:' + key), value)
                    borders.append(border)
                cell_props.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge, value in [('top', '65'), ('bottom', '65'), ('left', '95'), ('right', '95')]:
                    margin = OxmlElement('w:' + edge)
                    margin.set(qn('w:w'), value)
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                cell_props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.08
                    if column in numeric:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9.3)
                        run.bold = index == 0
                        run.font.color.rgb = RGBColor.from_string('FFFFFF' if index == 0 else '172A38')
        self.text.extend([' | '.join(headers)] + [' | '.join(map(str, row)) for row in rows])
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = .4

    def command(self, text):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(text)
        run.font.name = 'DejaVu Sans Mono'
        run.font.size = Pt(8.3)
        self.text.append(text)


def build_report(root):
    root = Path(root).resolve()
    receipt = load_receipts(root)
    summary, quality, audit = (receipt[k] for k in ('summary', 'quality', 'readiness'))
    rebuild, http, profile = (receipt[k] for k in ('rebuild', 'http', 'runtime'))
    counts = quality['counts']
    n = lambda value: f'{value:,}'
    output = root / 'reports'
    figures = output / 'figures'
    figures.mkdir(exist_ok=True)
    chart = figures / 'm2_snapshot_quality.png'
    draw_chart(summary, chart)
    report = Report()

    # Page 1: result, exact quantities and an interpretable visual.
    report.p('DS M01 到站時間與延誤預警', 'Title')
    report.p('M2 完整資料建置與驗證成果', 'Heading 1')
    report.p('Rollingtrans Data & AI Career Lab｜2026 年 10 月 27 日｜里程碑 v0.2.0', 'Caption')
    report.p(f"本輪已將 M1 試跑擴展為可保留及重建的 {n(summary['trip_count'])} 趟行程資料，"
             f"完成 {n(summary['snapshot_count'])} 張當時可知快照及對應標籤。全量資料品質、"
             '來源追溯、時間切分與 HTTP 對帳通過；這份資料可作為下一階段模型比較的固定輸入。')
    report.table(['項目', '本輪實測結果'], [
        ['資料期間', '2026 年 1 至 9 月；答案追蹤截止為 10 月 15 日 00:00 洛杉磯時間，不含該時點'],
        ['行程與定位', f"{n(summary['trip_count'])} 趟行程；{n(counts['telemetry_event'])} 筆 GPS 事件"],
        ['實體範圍', f"{counts['core_tenant']} 租戶、{counts['core_vehicle']} 車、{counts['core_device']} 裝置；"
                    f"{counts['eta_facility_version']} 場站版本、{counts['eta_route_version']} 路線版本"],
        ['歷史參考', f"{n(summary['history_reference_count'])} 筆可用訓練歷史關聯，完整保留而未抽樣截斷"],
        ['資料表與輸出', f"{quality['materialized_canonical_tables']} 張表結構及 Parquet；"
                      f"{quality['populated_canonical_tables']} 張有資料，{len(quality['reserved_empty_tables'])} 張待後續階段填入"],
    ], [1.25, 5.85])
    report.doc.add_picture(str(chart), width=Inches(7.1))
    report.p('圖 1　標籤資格與定位品質。單位是快照；同一趟行程可貢獻多張快照。', 'Caption')
    report.p('資料全部為明確標示的原創 2026 年合成情境，並非 Rollingtrans 客戶紀錄。'
             '正常與異常情境的比例屬工程設計，不能當作真實延誤率、預測準確度或營運改善。')

    # Page 2: independent sample support, role cutoffs and cohort cautions.
    report.page('1 時間切分與可評估樣本')
    report.p('訓練使用 1 月 1 日至 6 月 23 日出發的行程；7 月 1 至 24 日留給校準；'
             '8 至 9 月留給最終測試。訓練與校準答案分別只取 7 月 1 日、8 月 1 日之前可用的版本。'
             '緩衝期與測試前的保留車輛或路線行程均排除。')
    maturity = {(r['split_role'], r['label_eligibility']): r['snapshots']
                for r in audit['label_maturity_by_role']}
    roles = {r['split_role']: r for r in audit['trips_by_role']}
    report.table(['組別', '總行程', '可用目標行程', '可用目標快照', '未成熟快照'], [
        [ROLE_NAMES[role], n(roles[role]['trips']), n(roles[role]['eligible_trips']),
         n(roles[role]['eligible_snapshots']), n(maturity.get((role, 'not_mature'), 0))]
        for role in ('train', 'calibration', 'test', 'excluded')
    ], [.9, 1.1, 1.65, 1.7, 1.75], numeric=(1, 2, 3, 4))
    report.p('測試群覆蓋', 'Heading 2')
    report.table(['測試群', '總行程', '可用目標行程', '可用目標快照'], [
        [COHORT_NAMES[r['cohort']], n(r['trips']), n(r['eligible_trips']), n(r['eligible_snapshots'])]
        for r in audit['test_cohorts']
    ], [2.7, 1.1, 1.65, 1.65], numeric=(1, 2, 3))
    report.p('可用目標要求到場答案已在該組截止前確認，且快照時剩餘旅時為正。'
             '取消、未成熟與實際已到場的快照仍保留供稽核；不把未知填成零或準時。'
             '「已到場」也可能表示當時到場通知尚未送達，不能倒用事後資訊移除快照。'
             '排除組即使有合格標籤，也不納入訓練、校準或測試。')
    sparse = [r for r in audit['test_cohorts']
              if r['eligible_trips'] < audit['m3_support']['eligible_trip_diagnostic_floor']]
    if sparse:
        report.p('樣本限制：' + '；'.join(f"{COHORT_NAMES[r['cohort']]}僅 {r['eligible_trips']} 趟可用行程" for r in sparse)
                 + '。小群結果須分開呈現並揭露不確定性，不能只看快照總數。')
    no_history = sum(r['eligible_no_history_snapshots'] for r in audit['history_support_by_role'] if r['split_role'] == 'train')
    report.p(f'訓練組另有 {n(no_history)} 張可用目標快照沒有當時已知的訓練歷史池。M3 必須定義明確的'
             '無支援或備援策略；本輪的樣本門檻只是描述性提示，並非統計檢定力或上線驗收標準。')
    report.p('同趟快照彼此相關，模型驗證與不確定性估計須按行程分組。測試組本輪僅檢查資料有效性與覆蓋，'
             '不得據此選模型、特徵、校準方法或預警門檻。')

    # Page 3: actual receipts, engineering delta and schema/population distinction.
    report.page('2 驗證證據與工程變更')
    continuity = audit['m1_continuity']['files']
    report.table(['驗證項目', '實際結果'], [
        ['程式測試', f"{receipt['test_count']} 項測試全部通過；數量由本輪 unittest 執行紀錄讀取"],
        ['全量資料庫', f"{len(quality['checks'])} 項檢查通過，涵蓋型別、來源可用時間、標籤、歷史池與表結構"],
        ['建模前檢查', f"{len(audit['checks'])} 項技術檢查通過；群組樣本警示另行保留"],
        ['本機 HTTP', f"{n(http['pages'])} 頁、{n(http['output_rows'])} 筆事件；"
                     f"{http['retries']} 次重試，{http['duplicate_rows_replayed']} 筆重播正確去重"],
        ['乾淨目錄重建', f"{len(rebuild['files'])} 個資料檔 SHA256 相同；"
                       f"{len(rebuild['http_input_parity'])} 個 HTTP 衍生檔內容相同"],
        ['M1 程式相容', f"{len(receipt['compatibility']['files'])} 個 M1 衍生資料檔以優化後程式重建，逐位元組一致"],
        ['M1 資料延續', f"{len(continuity)} 個檔案內共 {n(sum(r['m1_row_count'] for r in continuity))} 筆 M1 業務及快照標籤資料，皆包含於 M2"],
    ], [1.45, 5.65])
    elapsed = float(profile['elapsed_seconds'])
    peak_gib = float(profile['peak_rss_kib']) / 1024 / 1024
    report.p(f'本次完整執行耗時 {elapsed:,.1f} 秒，量測到的程序最高常駐記憶體為 {peak_gib:.2f} GiB。'
             f"乾淨目錄重建及 HTTP 對帳另耗時 {float(rebuild['elapsed_seconds']):,.1f} 秒。"
             '這是本環境的實測紀錄，不代表其他電腦或正式服務的效能保證。')
    report.p('本輪工程調整', 'Heading 2')
    report.p('歷史關聯與來源追溯改為逐筆寫入，並以行程及租戶索引減少重複掃描；保留全部符合時間條件的歷史，'
             '沒有用抽樣縮減資料。輸出先寫暫存檔，完成並關閉後再逐檔發布，避免建置中途失敗覆蓋既有完整檔案。')
    report.p('M1 的生成器、種子、實體編號及特徵定義沿用。M2 取得時間與執行時間分開記錄；'
             '舊資料業務內容可對照，新的歷史池則隨全量訓練行程擴大。')
    report.p(f"{quality['materialized_canonical_tables']} 張必要表均已建立明確型別，其中 "
             f"{len(quality['reserved_empty_tables'])} 張模型、預測、輸出發布、預警通知與行動追蹤表為空表。"
             '表結構齊備不等於模型或預警流程完成；本輪沒有產生虛構預測或通知紀錄。')

    # Page 4: reproducible entry points and remaining work.
    report.page('3 重跑方式與 M3 入口')
    report.p('在解壓後的專案根目錄執行。需 Python 3.11 以上；已交付固定依賴版本與 M2 設定。'
             '重跑會更新生成資料與執行紀錄，請保留原始交付包作為參考快照。'
             '完整資料解壓約 12.6 GB；完整重建對帳建議預留約 30 GB 磁碟空間。')
    for command in (
        'python3 -m venv .venv',
        'source .venv/bin/activate',
        'python -m pip install -r requirements.lock',
        'python -m pip install -e . --no-deps',
        'python -m ds_m01.cli run --root . --config config/m2.json',
        'python -m ds_m01.cli http-check --root . --config config/m2.json',
        'python scripts/check_m2_readiness.py --root .',
        'python scripts/check_rebuild.py --root . --config config/m2.json',
        'python -m unittest discover -s tests -v',
    ):
        report.command(command)
    report.p('原生 PostgreSQL 狀態', 'Heading 2')
    pg = receipt['postgres']
    if pg.get('status') == 'passed':
        report.p('原生 PostgreSQL 往返測試已通過；具體資料表、列數與比對證據見 reports/postgres_check.json。')
    else:
        report.p('仍為未執行。此環境沒有可用的原生 PostgreSQL 服務、未設定資料庫連線，且僅映射 root 使用者；'
                 '不把 DuckDB 或本機 HTTP 通過視為 PostgreSQL 已驗證。DDL 與 psycopg 往返程式已提供，'
                 '可在相容環境設定 DS_M01_PG_DSN 後另行執行。')
    report.command('python -m ds_m01.postgres_check --root .')
    report.p('下一階段的模型工作', 'Heading 2')
    ready = audit.get('ready_to_attempt_m3') is True
    report.p(('資料技術檢查已支持開始 M3，仍須遵守小群與無歷史支援限制。' if ready else
              'M3 目前仍有樣本支援阻礙，應先處理 reports/m2_readiness.json 所列 blockers。')
             + '先建立歷史旅時基準，再比較剩餘旅時與分位數模型；只在訓練期內做時間驗證，'
             '用校準組處理區間、機率與預警門檻，最後一次性評估保留測試組。')
    report.p('尚待完成模型訓練、區間與機率校準、最終測試、預警通知與行動追蹤，以及完整專案驗收。'
             'GitHub 簽署 commit 與 push 由使用者操作；本輪未建立 commit、未 push。')
    report.p('主要證據位於 reports/quality_report.json、m2_readiness.json、rebuild_report.json、'
             'runtime_profile.json 與 unittest_results.txt；後續入口見 continuation.json。', 'Caption')
    report.doc.core_properties.title = 'DS M01 M2 完整資料建置與驗證成果'
    report.doc.core_properties.author = 'Rollingtrans Data & AI Career Lab'
    report.doc.save(output / (BASENAME + '.docx'))
    (output / (BASENAME + '.txt')).write_text('\n\n'.join(report.text) + '\n', encoding='utf-8')
    print(output / (BASENAME + '.docx'))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path('.'))
    build_report(parser.parse_args().root)
