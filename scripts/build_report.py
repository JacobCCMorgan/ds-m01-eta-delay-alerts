"""Build the Chinese milestone report and exact-count charts from verified results."""
from pathlib import Path
import argparse,json
from collections import Counter,defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path('.'));args=parser.parse_args()
root=args.root.resolve();out=root/'reports';fig=out/'figures';fig.mkdir(exist_ok=True)
read=lambda name:json.loads((out/name).read_text())
summary=read('m1_summary.json');quality=read('quality_report.json');rebuild=read('rebuild_report.json');http=read('last_http-check.json')
assert quality['passed'] and rebuild['status']=='passed' and http['status']=='passed'
config=json.loads((root/'config/m1.json').read_text())
generator=read('generator_validation.json');assert generator['status']=='passed'
def jsonl(p):return [json.loads(line) for line in p.read_text().splitlines() if line]
truth=jsonl(root/'validation/latent_truth.jsonl');snaps=jsonl(root/'data/processed/eta_snapshot.jsonl');labels=jsonl(root/'data/processed/eta_snapshot_label.jsonl')
promises=jsonl(root/'data/raw/eta_promise_version.jsonl')
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False})
figobj,axes=plt.subplots(1,2,figsize=(10.4,3.7),layout='constrained')
panels=[(['Eligible','Not mature','Cancelled','Already arrived'],[summary['label_eligibility'].get(k,0) for k in ['eligible','not_mature','cancelled','already_arrived']],'Snapshot target status'),
        (['Valid','Stale','No position'],[summary['position_quality'].get(k,0) for k in ['valid','low_confidence','no_position']],'Position availability')]
for ax,(names,values,title) in zip(axes,panels):
 ax.barh(names,values,color=['#226B88','#D89C42','#81909A','#AF6B60'][:len(names)]);ax.invert_yaxis();ax.set_title(title,loc='left',weight='bold');ax.set_xlabel('Snapshots');ax.set_xlim(0,max(values)*1.21)
 for i,v in enumerate(values):ax.text(v+5,i,str(v),va='center')
figobj.savefig(fig/'m1_snapshot_quality.png',dpi=180);plt.close(figobj)

case_names=['normal','late_arrival','promise_revision','gps_dropout','cancelled','arrival_report_late','outcome_revision','unfinished']
cases=[]
for scenario in case_names:
 trip=next(t for t in truth if t['scenario']==scenario)
 ss=sorted((s for s in snaps if s['trip_id']==trip['trip_id'] and s['tenant_id']==trip['tenant_id']),key=lambda x:x['snapshot_time'])
 selected={s['snapshot_id'] for s in ss}; ls=[l for l in labels if l['snapshot_id'] in selected]
 pp=sorted((p for p in promises if p['trip_id']==trip['trip_id'] and p['tenant_id']==trip['tenant_id']),key=lambda x:x['known_at'])
 cases.append({'scenario':scenario,'tenant_id':trip['tenant_id'],'trip_id':trip['trip_id'],'departure_at':trip['true_departure_at'],
               'first_snapshot_at':ss[0]['snapshot_time'] if ss else None,'initial_promise_end':pp[0]['window_end_at'],
               'actual_gate_arrival_at':trip['true_gate_arrival_at'],'gate_report_available_at':trip['gate_report_available_at'],
               'snapshot_count':len(ss),'label_counts':dict(Counter(l['label_eligibility'] for l in ls)),
               'position_counts':dict(Counter(s['position_validity'] for s in ss)),
               'evidence_is_synthetic':True,'hidden_physical_arrival_not_a_model_input':True})
(out/'representative_cases.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n')

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.78)
for n in ['Normal','Title','Heading 1','Heading 2','List Bullet','Caption']:
 st=doc.styles[n];st.font.name='Noto Sans CJK TC';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK TC');st.font.size=Pt(10.5);st.font.color.rgb=RGBColor.from_string('203344');st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 1'].font.color.rgb=RGBColor(0,0,0)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
doc.styles['Caption'].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('DS-M01 M1  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
text=[]
def p(s,style='Normal'):
 doc.add_paragraph(s,style);text.append(s)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E2ECF2' if ri==0 else ('F3F7FA' if ri%2 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
    for run in para.runs:run.font.size=Pt(9.5);run.bold=ri==0
 text.extend([' | '.join(headers)]+[' | '.join(map(str,row)) for row in rows]);doc.add_paragraph('')
def page(title):doc.add_page_break();p(title,'Heading 1')

p('DS-M01 到站時間與延誤預警','Title');p('M0 共用來源核對與 M1 資料試跑成果','Heading 1')
p('Rollingtrans Data & AI Career Lab｜2026 年 10 月 26 日｜里程碑 v0.1.0')
p('本輪已完成 120 趟合成行程的資料生成、當時可知快照、獨立到場標籤、HTTP 接收與 DuckDB／Parquet 驗證。這是資料工程與建模前置里程碑，尚未訓練 ETA 模型，也沒有可宣稱的預測準確度或營運改善。')
table(['項目','實際結果'],[['資料主期','2026/01/01–09/30；標籤追蹤截止為 10/15 00:00 洛杉磯時間，不含該時點'],['行程與實體','120 趟；3 租戶、30 車、36 裝置、12 場站及 12 路線'],['定位與快照',f"{quality['counts']['telemetry_event']:,} 筆定位；{summary['snapshot_count']:,} 張快照及標籤"],['資料契約','已啟用 17 張標準表；原必要表共 28 張，其餘 11 張保留到後續階段'],['驗證結果','33 項程式測試、24 項資料檢查及乾淨目錄重建通過'],['服務限制','本機 HTTP 已實際執行；原生 PostgreSQL 往返測試尚未執行']],[1.3,5.25])
p('全部行程與定位皆為原創模擬資料；不是 Rollingtrans 客戶記錄。情境刻意平均配置正常與異常案例，比例不能當作真實車隊的延誤率或資料品質。')
p('本輪沒有建立 Git commit 或 push。程式、資料、操作文件與驗證證據全部交付，由使用者自行建立簽署 commit。DA-E01 的原始檔案保持不變，仍待獨立驗收。')

page('1 資料結果與可評估範圍')
doc.add_picture(str(fig/'m1_snapshot_quality.png'),width=Inches(6.55));p('圖 1　快照的目標資格與定位品質。單位是快照，不是獨立行程。','Caption')
table(['快照標籤','數量','解讀'],[['可用目標',457,'到場答案已在相應截止前成熟，剩餘時間為正'],['尚無成熟答案',288,'待確認或未完成；保留未知，不能填零或準時'],['取消',11,'保留快照及取消結果，排除剩餘旅時目標'],['已到場',181,'實際已到場但當時系統未得知，保留快照並排除迴歸目標']],[1.3,.6,4.65])
p('定位有效 413 張、陳舊而低可信度 465 張、無可用定位 59 張。無可靠定位時，不推算剩餘距離。較多陳舊快照主要來自刻意設計的晚到場通知與追蹤中斷情境。')
p('行程按時間與保留群分為訓練 65、校準 13、測試 26、排除 16 趟。同趟快照固定同組；未見車輛與路線不得進入訓練、校準或歷史旅時池。這些分組已凍結，但本輪沒有用它們訓練或調整模型。')
p('每次快照最長追蹤出發後 12 小時。這是 M1 的明示觀測上限，不代表未完成行程已到場；正式母體、標籤成熟與跨界追蹤仍須在 M2 完整資料中檢查。')

page('2 代表案例與真實修正紀錄')
descriptions={
 'normal':'一般來源與到場流程；對照最初承諾及到場標籤。',
 'late_arrival':'以當時承諾計算延誤真值；本輪沒有產生延誤機率。',
 'promise_revision':'保留修改前後承諾版本，快照只使用當時已知版本。',
 'gps_dropout':'定位中斷期間沿用最後觀測並揭露陳舊，不捏造位置。',
 'cancelled':'出發後 25 分鐘取消，保留收到出發後的第一張快照。',
 'arrival_report_late':'到場訊息故意晚 20 天送達，檢查事後答案不回寫特徵。',
 'outcome_revision':'先到的答案偏早 5 分鐘，後續修訂依可用時間選取。',
 'unfinished':'來源追蹤中斷；隱藏物理行程結束不等於已有可用標籤。'}
table(['情境與案例 ID','本輪檢查'],[[f"{c['scenario']}\n{c['trip_id']}",descriptions[c['scenario']]] for c in cases],[2.3,4.25])
p('完整時間線、快照數與標籤分布見 reports/representative_cases.json。該檔供查閱及驗證，包含模擬真值，不可加入模型特徵。')
p('實際修正','Heading 2')
p('第一張快照改在收到出發事件時建立，避免短程取消被漏掉；未知或撤回的裝置關聯不再視為有效；歷史旅時改採實際出發時間；衍生資料可用時間不得早於任何來源；後續狀態更正也納入來源追溯。')

page('3 重現證據與未完成事項')
table(['檢查','證據與範圍'],[['M0 共用來源','原 ZIP SHA256 已核對，10 組窄相容性檢查通過；165 個 DA-E01 檔案內容前後相同'],['獨立程式測試','20 項時間邏輯、4 項建置整合、8 項 HTTP、1 項來源完整性，共 33 項'],['資料庫品質','24 項檢查涵蓋來源截止、標籤真值、歷史池、來源關聯與型別'],['實際 HTTP','94 頁；一次刻意注入 503 後成功重試；重播 137 筆後正確去重，內容一致'],['乾淨目錄重建',f"{len(rebuild['files'])} 個資料檔逐一 SHA256 相符；HTTP 接收資料產生的衍生結果亦一致"],['擴量穩定性','暫存建立 120／1,800 趟，原 120 趟業務資料完全包含於全量設計，排程及裝置關聯檢查通過']],[1.3,5.25])
p('原生 PostgreSQL 的 DDL 與 psycopg 往返檢查程式已提供。此環境沒有原生服務且無法切換一般使用者啟動，因此未宣稱 PostgreSQL 實測通過；完整服務模式對帳仍是下一工程項目。HTTP 通過只涵蓋定位接收及與營運快照的結果一致性。')
p('尚未完成：保留完整 1,800 趟端到端資料、其餘 11 張必要表、基準與 ETA 模型、區間／機率校準、最終模型評估、預警通知與行動追蹤、使用者實機重跑及驗收。')
p('M1 以固定的原始生成時間戳重播資料，另記錄每次指令的實際執行時間。情境日期是合成歷史，不能解讀成模型在 2026 年年初已部署。')

page('4 操作入口與下一階段')
p('建議先閱讀 README.zh-TW.md。Mac 上需 Python 3.11 以上；本輪實測 Linux／Python 3.12.14，未冒稱完成 Mac 實機驗證。資料重跑不需要公司檔案、外部 API key 或付費服務。')
table(['入口','用途'],[['python -m ds_m01.cli run --root .','生成 120 趟、快照與標籤、DuckDB／Parquet 及品質檢查'],['python -m ds_m01.cli http-check --root .','啟動暫時的本機 HTTP 服務並驗證接收'],['python -m unittest discover -s tests -v','執行獨立邊界與整合測試'],['python scripts/check_rebuild.py --root .','乾淨目錄重建及 HTTP 接收結果對帳'],['reports/quality_report.json','最新資料檢查結果；任何失敗都應先排查'],['continuation.json','本輪完成範圍、限制與下一個工作入口']],[3.5,3.05])
p('下一階段先處理 PostgreSQL 實際服務驗證，再將已驗證的生成器擴至 1,800 趟並重跑全量資料品質、切分與標籤成熟檢查；之後才開始模型比較與校準。若服務環境仍不可用，會保留阻礙，並分開記錄可繼續完成的資料工作。')
p('目前不需要使用者再批准例行程式與資料工作。使用者可以針對本報告的代表情境、結果說明及操作可讀性提出調整；GitHub 的簽署 commit 與 push 由使用者處理。')
p('English summary','Heading 2')
p('M1 delivers a reproducible synthetic ETA data pipeline, not a trained forecasting model. The 120-trip trial produces 937 point-in-time snapshots across 17 canonical tables. Independent arrival evidence, time-aware joins, immutable promises, and restricted historical labels are checked. Local HTTP ingestion and fresh rebuild parity pass; native PostgreSQL execution remains unverified. No business impact or real-fleet accuracy is claimed.')
doc.core_properties.title='DS-M01 M0 與 M1 試跑成果';doc.core_properties.author='Rollingtrans Data & AI Career Lab'
name='DS-M01_M1_Results_20261026_v0_1'
doc.save(out/(name+'.docx'));(out/(name+'.txt')).write_text('\n\n'.join(text)+'\n')
print(out/(name+'.docx'))
